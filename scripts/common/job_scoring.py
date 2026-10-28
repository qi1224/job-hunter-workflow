"""Explainable profile-driven scoring; no candidate-specific defaults."""
import re
from pathlib import Path

def load_resume_text(resume_path):
    path = Path(resume_path).expanduser()
    if path.suffix.lower() == '.docx':
        from docx import Document
        doc = Document(str(path))
        return '\n'.join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    if path.suffix.lower() == '.pdf':
        from PyPDF2 import PdfReader
        return '\n'.join(p.extract_text() or '' for p in PdfReader(str(path)).pages)
    if path.suffix.lower() not in ('.txt', '.md'):
        raise ValueError('Use PDF, DOCX, TXT or Markdown CV; scanned PDFs need OCR first')
    return path.read_text(encoding='utf-8-sig')

def contains(term, text):
    return bool(re.search(r'(?<!\w)' + re.escape(term) + r'(?!\w)', text, re.I))

def score_jobs(jobs, resume_path=None, preferred_countries=None, profile=None):
    if not profile:
        raise ValueError('Missing scoring profile. Complete CV onboarding first.')
    cv = load_resume_text(resume_path) if resume_path else ''
    skills = [s for s in profile['skills'] if contains(s['evidence'], cv)]
    if not skills:
        raise ValueError('No profile evidence matches CV; repeat onboarding')
    denominator = sum(s['weight'] for s in skills)
    result = []
    for original in jobs:
        job = dict(original)
        title = job.get('title') or ''
        text = title + ' ' + (job.get('description') or '')
        required = profile.get('required_any', [])
        if required and not any(contains(t, text) for t in required):
            continue
        matched = [s for s in skills if any(contains(t, text) for t in s['terms'])]
        skill_score = round(75 * sum(s['weight'] for s in matched) / denominator)
        title_score = 25 if any(contains(t, title) for t in profile['title_terms']) else 0
        excluded = [t for t in profile.get('excluded_title_terms', []) if contains(t, title)]
        job['match_score'] = max(0, skill_score + title_score - (35 if excluded else 0))
        job['match_reason'] = (f'skills {skill_score}/75 + title {title_score}/25; matched: ' +
                              ', '.join(s['label'] for s in matched) +
                              ('; excluded title -35: ' + ', '.join(excluded) if excluded else ''))
        result.append(job)
    return sorted(result, key=lambda j: j['match_score'], reverse=True)
